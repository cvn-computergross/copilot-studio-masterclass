"""Genera il Dossier di revisione di una proposta di progetto IT (Word) con un layout fisso.

Uso:
    python generate_dossier.py dossier_input.json dossier_revisione.docx

Contratto JSON di input (tutti i campi sono obbligatori salvo diversa indicazione):
{
  "codice_valutazione": "AS-1001",
  "proposta": {
    "codice": "PP-1001",
    "titolo": "Portale fornitori self-service",
    "sintesi": "Paragrafo di sintesi per i revisori",
    "budget_richiesto": 180000,
    "durata_mesi": 9
  },
  "proponente": {"nome": "Giulia Ferrari", "ruolo": "Responsabile Ufficio Acquisti",
                 "funzione": "Acquisti", "email": "giulia.ferrari@zava-demo.example"},
  "programma": {"codice": "PG-100", "nome": "Digital Process Transformation",
                "budget_disponibile": 600000},
  "criteri": [
    {"nome": "Strategic Alignment", "peso": 25, "livello": "Forte",
     "evidenza": "Frase breve tratta dal business case"}
  ],
  "domande": [
    {"criterio": "Strategic Alignment", "domanda": "Testo della domanda",
     "verifica": "Che cosa verifica la domanda"}
  ],
  "rischi": ["Rischio o punto di attenzione"],          (facoltativo)
  "punteggio_registrato": 87.0                          (facoltativo: csm_score della Valutazione)
}

Regole applicate dal codice (non dal modello):
- i livelli ammessi sono Forte (1.0), Moderata (0.6), Debole (0.3), Assente (0.0);
- i pesi dei criteri devono sommare a 100;
- punti = fattore x peso; punteggio complessivo = somma dei punti;
- le domande devono essere esattamente 8, ciascuna collegata a un criterio esistente,
  e vengono raggruppate per criterio a partire dal peso più alto;
- se il budget richiesto supera il budget disponibile del programma, viene aggiunto un rischio;
- se e' presente punteggio_registrato e differisce dal punteggio calcolato di oltre 0,1 punti,
  viene aggiunto un punto di attenzione (il dossier deve riflettere la valutazione registrata).
"""
import datetime as dt
import json
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

FATTORI = {"Forte": 1.0, "Moderata": 0.6, "Debole": 0.3, "Assente": 0.0}
NUMERO_DOMANDE = 8


class DossierError(ValueError):
    """Errore di validazione dei dati di input."""


def euro(value):
    s = f"{float(value):,.2f}"
    return "€ " + s.replace(",", "X").replace(".", ",").replace("X", ".")


def numero(value):
    return f"{value:.1f}".replace(".", ",")


def valida(data):
    for key in ("codice_valutazione", "proposta", "proponente", "programma", "criteri", "domande"):
        if key not in data:
            raise DossierError(f"Campo obbligatorio mancante: {key}")
    criteri = data["criteri"]
    if not criteri:
        raise DossierError("La lista dei criteri è vuota")
    for c in criteri:
        if c.get("livello") not in FATTORI:
            raise DossierError(f"Livello non ammesso per '{c.get('nome')}': {c.get('livello')}. "
                               f"Ammessi: {', '.join(FATTORI)}")
        if not str(c.get("evidenza", "")).strip():
            raise DossierError(f"Evidenza mancante per il criterio '{c.get('nome')}'")
    totale_pesi = sum(int(c["peso"]) for c in criteri)
    if totale_pesi != 100:
        raise DossierError(f"I pesi dei criteri sommano a {totale_pesi}, non a 100")
    nomi = {c["nome"] for c in criteri}
    domande = data["domande"]
    if len(domande) != NUMERO_DOMANDE:
        raise DossierError(f"Servono esattamente {NUMERO_DOMANDE} domande, ricevute {len(domande)}")
    for q in domande:
        if q.get("criterio") not in nomi:
            raise DossierError(f"La domanda '{q.get('domanda')}' non è collegata a un criterio esistente")


def calcola(data):
    righe = []
    for c in data["criteri"]:
        punti = round(FATTORI[c["livello"]] * int(c["peso"]), 1)
        righe.append({**c, "peso": int(c["peso"]), "punti": punti})
    righe.sort(key=lambda r: -r["peso"])
    totale = round(sum(r["punti"] for r in righe), 1)
    ordine = {r["nome"]: i for i, r in enumerate(righe)}
    domande = sorted(data["domande"], key=lambda q: ordine[q["criterio"]])
    rischi = list(data.get("rischi", []))
    budget = float(data["proposta"]["budget_richiesto"])
    disponibile = float(data["programma"]["budget_disponibile"])
    if budget > disponibile:
        rischi.insert(0, f"Il budget richiesto ({euro(budget)}) supera il budget disponibile del programma "
                         f"({euro(disponibile)}).")
    registrato = data.get("punteggio_registrato")
    if registrato is not None and abs(float(registrato) - totale) > 0.1:
        rischi.append(f"Il punteggio calcolato nel dossier ({numero(totale)}) differisce da quello registrato "
                      f"nella valutazione ({numero(float(registrato))}): verificare i livelli prima della revisione.")
    return righe, totale, domande, rischi


def run(doc, text, size=10.5, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if align is not None:
        p.alignment = align
    return p


def bullet(doc, label, value):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(label).bold = True
    p.add_run(value)


def genera(data, output):
    valida(data)
    righe, totale, domande, rischi = calcola(data)
    pr, pp, pg = data["proposta"], data["proponente"], data["programma"]

    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    run(d, "Dossier di revisione della proposta", 20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(d, pr["titolo"], 12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    oggi = dt.date.today().strftime("%d/%m/%Y")
    run(d, f"Proposta {pr['codice']} - Valutazione {data['codice_valutazione']} - Generato il {oggi}", 9,
        align=WD_ALIGN_PARAGRAPH.CENTER)

    d.add_heading("Proponente", level=2)
    bullet(d, "Nome: ", pp["nome"])
    bullet(d, "Ruolo e funzione: ", f"{pp['ruolo']} - {pp['funzione']}")
    bullet(d, "Email: ", pp["email"])

    d.add_heading("Programma", level=2)
    bullet(d, "Programma: ", f"{pg['codice']} {pg['nome']}")
    bullet(d, "Budget richiesto: ", f"{euro(pr['budget_richiesto'])} (durata {pr['durata_mesi']} mesi)")
    bullet(d, "Budget disponibile del programma: ", euro(pg["budget_disponibile"]))

    d.add_heading("Sintesi della proposta", level=2)
    run(d, pr["sintesi"])

    d.add_heading("Criteri di valutazione ed evidenze", level=2)
    t = d.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(["Criterio", "Peso", "Livello", "Evidenza", "Punti"]):
        t.rows[0].cells[i].text = h
    for r in righe:
        cells = t.add_row().cells
        cells[0].text = r["nome"]
        cells[1].text = f"{r['peso']}%"
        cells[2].text = r["livello"]
        cells[3].text = r["evidenza"]
        cells[4].text = numero(r["punti"])
    p = run(d, f"Punteggio complessivo: {numero(totale)}/100", 11, bold=True)
    p.paragraph_format.space_before = Pt(6)

    d.add_heading("Domande per lo sponsor", level=2)
    corrente = None
    pesi = {r["nome"]: r["peso"] for r in righe}
    for n, q in enumerate(domande, start=1):
        if q["criterio"] != corrente:
            corrente = q["criterio"]
            run(d, f"{corrente} (peso {pesi[corrente]}%)", 10.5, bold=True)
        d.add_paragraph(f"{n}. {q['domanda']}")
        if q.get("verifica"):
            run(d, f"    Verifica: {q['verifica']}", 9.5, italic=True)

    d.add_heading("Rischi e punti di attenzione", level=2)
    if rischi:
        for r in rischi:
            d.add_paragraph(r, style="List Bullet")
    else:
        run(d, "Nessun rischio specifico segnalato.")

    run(d, "Documento di supporto preparato con l'intelligenza artificiale. Verificare i contenuti prima "
           "dell'uso. Livelli e punteggi sono proposte: la decisione sulla proposta spetta al Comitato di "
           "Valutazione.", 8.5, italic=True)
    d.save(output)
    return {"file": output, "punteggio": totale, "criteri": len(righe), "domande": len(domande)}


def main(argv):
    if len(argv) != 3:
        print("Uso: python generate_dossier.py <input.json> <output.docx>", file=sys.stderr)
        return 2
    with open(argv[1], encoding="utf-8") as f:
        data = json.load(f)
    try:
        esito = genera(data, argv[2])
    except DossierError as e:
        print(f"ERRORE DI VALIDAZIONE: {e}", file=sys.stderr)
        return 1
    print(json.dumps(esito, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
